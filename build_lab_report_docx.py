"""Build professional Lab 04 Word (.docx) report with embedded terminal screenshots."""

import os

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

DOC_NAME = "Lab_04_Report_Zaeem_Abbas.docx"
SCREENSHOTS_DIR = os.path.abspath("screenshots")

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)

    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    # Add border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="D1D5DB"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="D1D5DB"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="D1D5DB"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15

    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(31, 41, 55)

    # Empty paragraph after code block
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def build_report():
    doc = docx.Document()

    # Set page margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("LAB 04 REPORT")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(17, 24, 39)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Testing with Pytest: Fixtures, Parametrization, and Test Doubles")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(75, 85, 99)

    # Student Info Table
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.columns[0].width = Inches(2.2)
    info_table.columns[1].width = Inches(4.3)

    metadata = [
        ("Student Name:", "Zaeem Abbas Salim Abbas"),
        ("Roll Number:", "2430-184"),
        ("Program & Semester:", "BS Software Engineering (Semester 5)"),
        ("Course:", "Software Construction and Design Lab"),
        ("Repository Link:", "https://github.com/abbaszaeem89-bit/gradebook-service.git"),
    ]

    for i, (label, val) in enumerate(metadata):
        row = info_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        set_cell_background(c1, "F3F4F6")
        set_cell_background(c2, "FFFFFF")
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c2, top=80, bottom=80, left=120, right=120)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(31, 41, 55)

        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(val)
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(17, 24, 39)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Tasks Content Definition
    tasks = [
        {
            "num": 1,
            "title": "Task 1: Pytest Configuration (pytest.ini)",
            "desc": "Configured pytest.ini at the project root directory. The setting pythonpath = . allows test files to import project modules (gradebook, models, errors) directly without manual sys.path manipulation, while testpaths = tests instructs pytest to discover tests inside the tests/ directory.",
            "code_label": "File: pytest.ini",
            "code": """[pytest]
pythonpath = .
testpaths = tests""",
            "screenshot": "task1_terminal.png",
            "output_desc": "Terminal output verifying that pytest successfully loaded pytest.ini, located the tests directory, and passed all existing test suites.",
        },
        {
            "num": 2,
            "title": "Task 2: Initial Tests for pct() using monkeypatch",
            "desc": "Constructed tests/test_pct.py using the Arrange-Act-Assert pattern. Because pct() reads global STATE, monkeypatch.setattr is used to inject an isolated dictionary for each test and automatically revert state afterwards. Tested three scenarios: a student with no marks, standard weight percentage calculation (15/20 on weight 10 = 7.5%), and ignoring marks for unknown assessments.",
            "code_label": "File: tests/test_pct.py (Initial Version)",
            "code": """import gradebook


def test_student_with_no_marks_has_zero_percent(monkeypatch):
    state = {
        "students": {"S1": {"id": "S1", "name": "Ayesha"}},
        "assessments": {"A1": {"id": "A1", "title": "Quiz 1", "weight": 10, "total": 20}},
        "marks": [],
    }
    monkeypatch.setattr(gradebook, "STATE", state)
    result = gradebook.pct("S1")
    assert result == "0"


def test_mark_of_15_out_of_20_on_weight_10_gives_7_5(monkeypatch):
    state = {
        "students": {"S1": {"id": "S1", "name": "Ayesha"}},
        "assessments": {"A1": {"id": "A1", "title": "Quiz 1", "weight": 10, "total": 20}},
        "marks": [{"student": "S1", "assessment": "A1", "score": 15}],
    }
    monkeypatch.setattr(gradebook, "STATE", state)
    result = gradebook.pct("S1")
    assert result == "7.5"


def test_mark_for_unknown_assessment_is_ignored(monkeypatch):
    state = {
        "students": {"S1": {"id": "S1", "name": "Ayesha"}},
        "assessments": {},
        "marks": [{"student": "S1", "assessment": "GHOST", "score": 15}],
    }
    monkeypatch.setattr(gradebook, "STATE", state)
    result = gradebook.pct("S1")
    assert result == "0" """,
            "screenshot": "task2_terminal.png",
            "output_desc": "Terminal output executing pytest tests/test_pct.py -v, confirming 3 passed tests.",
        },
        {
            "num": 3,
            "title": "Task 3: Shared Fixtures in conftest.py",
            "desc": "Extracted duplicated test setup into pytest fixtures inside tests/conftest.py. The state fixture builds a fresh mock gradebook dictionary, and the loaded fixture monkeypatches gradebook.STATE and returns it. Each test receives a clean, isolated copy, eliminating test pollution.",
            "code_label": "Files: tests/conftest.py & Refactored tests/test_pct.py",
            "code": """# tests/conftest.py
from typing import Any
import pytest
import gradebook

@pytest.fixture
def state() -> dict[str, Any]:
    return {
        "students": {"S1": {"id": "S1", "name": "Ayesha"}},
        "assessments": {
            "A1": {"id": "A1", "title": "Quiz 1", "weight": 10, "total": 20},
            "A2": {"id": "A2", "title": "Midterm", "weight": 30, "total": 50},
        },
        "marks": [],
    }

@pytest.fixture
def loaded(state: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    monkeypatch.setattr(gradebook, "STATE", state)
    return state

# tests/test_pct.py (Refactored to use loaded fixture)
import gradebook

def test_student_with_no_marks_has_zero_percent(loaded):
    result = gradebook.pct("S1")
    assert result == "0"

def test_mark_of_15_out_of_20_on_weight_10_gives_7_5(loaded):
    loaded["marks"].append({"student": "S1", "assessment": "A1", "score": 15})
    result = gradebook.pct("S1")
    assert result == "7.5"

def test_mark_for_unknown_assessment_is_ignored(loaded):
    loaded["marks"].append({"student": "S1", "assessment": "GHOST", "score": 15})
    result = gradebook.pct("S1")
    assert result == "0" """,
            "screenshot": "task3_terminal.png",
            "output_desc": "Terminal output verifying that tests/test_pct.py runs cleanly against the shared fixtures.",
        },
        {
            "num": 4,
            "title": "Task 4: Parametrizing Tests across Multiple Scenarios",
            "desc": "Consolidated repetitive test cases into @pytest.mark.parametrize tables. In tests/test_pct.py, five score scenarios (no marks, partial quiz, full marks, zero score, two combined assessments) are tested in a single clean function. In tests/test_models.py, numeric parsing is parametrized into good numbers accepted and invalid types/tokens rejected.",
            "code_label": "Files: tests/test_pct.py & Parametrized tests/test_models.py",
            "code": """# tests/test_pct.py
import pytest
import gradebook

@pytest.mark.parametrize(
    ("scores", "expected"),
    [
        pytest.param({}, "0", id="no-marks"),
        pytest.param({"A1": 15}, "7.5", id="quiz-15-of-20"),
        pytest.param({"A1": 20}, "10.0", id="quiz-full-marks"),
        pytest.param({"A1": 0}, "0.0", id="quiz-zero-score"),
        pytest.param({"A1": 15, "A2": 40}, "31.5", id="two-assessments"),
    ],
)
def test_percentage_for_scores(loaded, scores, expected):
    for assessment, score in scores.items():
        loaded["marks"].append({"student": "S1", "assessment": assessment, "score": score})
    result = gradebook.pct("S1")
    assert result == expected

def test_mark_for_unknown_assessment_is_ignored(loaded):
    loaded["marks"].append({"student": "S1", "assessment": "GHOST", "score": 15})
    result = gradebook.pct("S1")
    assert result == "0"

# tests/test_models.py (Parametrized Number Tests)
@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        pytest.param("15", 15.0, id="numeric-string"),
        pytest.param(15, 15.0, id="int"),
        pytest.param("2.5", 2.5, id="decimal-string"),
    ],
)
def test_good_numbers_are_accepted(raw: object, expected: float) -> None:
    assert parse_number(raw, "score") == expected

@pytest.mark.parametrize(
    "raw",
    [
        pytest.param("abc", id="text"),
        pytest.param(True, id="bool"),
        pytest.param(None, id="none"),
        pytest.param([], id="list"),
        pytest.param("nan", id="nan"),
        pytest.param("inf", id="inf"),
    ],
)
def test_bad_numbers_are_rejected(raw: object) -> None:
    with pytest.raises(ValidationError, match="score"):
        parse_number(raw, "score") """,
            "screenshot": "task4_terminal.png",
            "output_desc": "Terminal output executing pytest -v, showing all 20 parametrized test cases passing.",
        },
        {
            "num": 5,
            "title": "Task 5: Testing File Persistence with tmp_path",
            "desc": "Created tests/test_storage.py to test load() and save() against real temporary files using pytest's built-in tmp_path fixture. Avoiding mocks ensures that actual disk I/O, serialization, missing-file defaults, and JSON corruption errors (json.JSONDecodeError) are genuinely verified.",
            "code_label": "File: tests/test_storage.py",
            "code": """import json
from pathlib import Path
from typing import Any
import pytest
import gradebook

@pytest.fixture
def data_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    path = tmp_path / "gradebook.json"
    monkeypatch.setattr(gradebook, "DATA", str(path))
    return path

def test_saved_gradebook_loads_back_unchanged(data_file, state, monkeypatch):
    monkeypatch.setattr(gradebook, "STATE", state)
    gradebook.save()
    monkeypatch.setattr(gradebook, "STATE", {})
    gradebook.load()
    assert gradebook.STATE == state

def test_missing_file_gives_an_empty_gradebook(data_file, monkeypatch):
    monkeypatch.setattr(gradebook, "STATE", {})
    gradebook.load()
    assert gradebook.STATE == {"students": {}, "assessments": {}, "marks": []}

def test_corrupt_file_stops_loading_with_a_json_error(data_file):
    data_file.write_text("{bad")
    with pytest.raises(json.JSONDecodeError):
        gradebook.load() """,
            "screenshot": "task5_terminal.png",
            "output_desc": "Terminal output executing pytest tests/test_storage.py -v, showing all 3 storage tests passing (23 total).",
        },
        {
            "num": 6,
            "title": "Task 6: Creating an Architectural Seam with add_mark()",
            "desc": "Extracted mark validation, recording, timestamp generation, and persistence out of the HTTP handler into a pure domain function add_mark(state, mark, now, save). By passing now (clock) and save (persistence) as parameters, the function can be tested in complete isolation without launching an HTTP server.",
            "code_label": "File: gradebook.py (Seam Extraction)",
            "code": """def now() -> str:
    return str(datetime.datetime.now(datetime.timezone.utc))

def add_mark(
    state: dict[str, Any],
    mark: Mark,
    now: Callable[[], str],
    save: Callable[[], None],
) -> None:
    \"\"\"Check a mark against the gradebook, then record it and save.\"\"\"
    if mark.student not in state["students"]:
        raise NotFoundError(f"no such student: {mark.student}")
    assessment_row = state["assessments"].get(mark.assessment)
    if assessment_row is None:
        raise NotFoundError(f"no such assessment: {mark.assessment}")
    if mark.score > assessment_row["total"]:
        raise ValidationError("score exceeds assessment total")
    state["marks"].append(
        {
            "student": mark.student,
            "assessment": mark.assessment,
            "score": mark.score,
            "at": now(),
        }
    )
    save()

# Handler integration inside do_POST:
if self.path == "/marks":
    mark = parse_mark(payload)
    add_mark(STATE, mark, now, save)
    logger.info("recorded mark for %s score %s", mark.student, mark.score)
    return 200, {"ok": True} """,
            "screenshot": "task6_terminal.png",
            "output_desc": "Terminal output verifying ruff check ., mypy gradebook.py, and pytest with 0 issues.",
        },
        {
            "num": 7,
            "title": "Task 7: Testing add_mark with Test Doubles (Stub, Fake, Mock)",
            "desc": "Constructed tests/test_add_mark.py using three distinct test doubles: a Stub (fixed_now providing a deterministic ISO time string), a Fake (state fixture acting as in-memory storage), and a Mock (unittest.mock.Mock standing in for save to verify call counts). Verified that rejected marks leave state completely clean and never invoke save().",
            "code_label": "File: tests/test_add_mark.py",
            "code": """from typing import Any
from unittest.mock import Mock
import pytest
from errors import NotFoundError, ValidationError
from gradebook import add_mark
from models import Mark

def fixed_now() -> str:
    return "2026-10-12 09:00:00"

def test_valid_mark_is_stored_with_the_time(state: dict[str, Any]) -> None:
    save = Mock()
    add_mark(state, Mark("S1", "A1", 15.0), fixed_now, save)
    assert state["marks"] == [
        {"student": "S1", "assessment": "A1", "score": 15.0, "at": "2026-10-12 09:00:00"}
    ]

def test_valid_mark_calls_save_once(state: dict[str, Any]) -> None:
    save = Mock()
    add_mark(state, Mark("S1", "A1", 15.0), fixed_now, save)
    save.assert_called_once_with()

@pytest.mark.parametrize(
    ("mark", "error"),
    [
        pytest.param(Mark("NOPE", "A1", 5.0), NotFoundError, id="unknown-student"),
        pytest.param(Mark("S1", "GHOST", 5.0), NotFoundError, id="unknown-assessment"),
        pytest.param(Mark("S1", "A1", 21.0), ValidationError, id="score-above-total"),
    ],
)
def test_rejected_mark_changes_nothing(state: dict[str, Any], mark: Mark, error: type[Exception]) -> None:
    save = Mock()
    with pytest.raises(error):
        add_mark(state, mark, fixed_now, save)
    assert state["marks"] == []
    save.assert_not_called() """,
            "screenshot": "task7_terminal.png",
            "output_desc": "Terminal output executing pytest tests/test_add_mark.py -v, confirming 5 passed tests (28 passed in total).",
        },
        {
            "num": 8,
            "title": "Task 8: Fault Injection & Proving Tests Can Fail",
            "desc": "Conducted fault injection to prove test suite efficacy. Temporarily altered division (/) to multiplication (*) in pct(). Pytest immediately detected the flaw, resulting in '3 failed, 25 passed' across all non-zero percentage cases. After reverting the mutation, all 28 tests passed cleanly.",
            "code_label": "Mutation in gradebook.py & NOTES.md Section",
            "code": """# gradebook.py (Temporary Fault Injection)
# Original: got = got + float(m["score"]) * float(a["weight"]) / float(a["total"])
got = got + float(m["score"]) * float(a["weight"]) * float(a["total"])  # Broken on purpose

# Result: 3 failed, 25 passed

# After Reverting:
got = got + float(m["score"]) * float(a["weight"]) / float(a["total"])  # Clean
# Result: 28 passed

# NOTES.md (Excerpts):
# Which double is which:
# - Stub: fixed_now in tests/test_add_mark.py (deterministic return value)
# - Fake: state fixture in tests/conftest.py (working in-memory gradebook)
# - Mock: Mock() passed as save in tests/test_add_mark.py (records interactions)
#
# Why pct() and files were not mocked:
# - pct() is pure internal business logic; mocking it would bypass the calculation under test.
# - tmp_path provides a real, isolated file system to genuinely verify JSON disk read/write. """,
            "screenshot": "task8_terminal.png",
            "output_desc": "Terminal output showing the fault injection failure (3 failed, 25 passed) followed by the successful re-run (28 passed).",
        },
    ]

    for t in tasks:
        # Heading
        h = doc.add_heading(t["title"], level=1)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.name = "Calibri"
            r.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy

        # Description
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_after = Pt(6)
        p_desc.paragraph_format.line_spacing = 1.15
        r_desc = p_desc.add_run(t["desc"])
        r_desc.font.name = "Calibri"
        r_desc.font.size = Pt(10.5)

        # Code Label
        p_lbl = doc.add_paragraph()
        p_lbl.paragraph_format.space_before = Pt(4)
        p_lbl.paragraph_format.space_after = Pt(2)
        r_lbl = p_lbl.add_run(t["code_label"])
        r_lbl.font.bold = True
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = RGBColor(55, 65, 81)

        # Code Block
        add_code_block(doc, t["code"])

        # Terminal Output description
        p_out = doc.add_paragraph()
        p_out.paragraph_format.space_before = Pt(4)
        p_out.paragraph_format.space_after = Pt(4)
        r_out_title = p_out.add_run("Terminal Output Verification:\n")
        r_out_title.font.bold = True
        r_out_title.font.size = Pt(10)
        r_out_title.font.color.rgb = RGBColor(55, 65, 81)

        r_out_desc = p_out.add_run(t["output_desc"])
        r_out_desc.font.italic = True
        r_out_desc.font.size = Pt(9.5)
        r_out_desc.font.color.rgb = RGBColor(107, 114, 128)

        # Screenshot
        img_path = os.path.join(SCREENSHOTS_DIR, t["screenshot"])
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(4)
            p_img.paragraph_format.space_after = Pt(12)
            doc.add_picture(img_path, width=Inches(6.5))

    # Final Summary & Checklist
    h_sum = doc.add_heading("Final Quality & Verification Summary", level=1)
    h_sum.paragraph_format.space_before = Pt(16)
    h_sum.paragraph_format.space_after = Pt(6)
    for r in h_sum.runs:
        r.font.name = "Calibri"
        r.font.color.rgb = RGBColor(30, 58, 138)

    p_check = doc.add_paragraph()
    p_check.paragraph_format.line_spacing = 1.2
    checklist_items = [
        "All 28 unit tests passing with pytest (pytest tests/).",
        "Ruff linter reports zero errors across all modules and tests (ruff check .).",
        "Mypy static type checker passes without issues (mypy gradebook.py).",
        "Strict static type checking passes for domain models and errors (mypy --strict errors.py models.py).",
        "Seam add_mark() successfully decouples business logic from HTTP transport layer.",
        "Deterministic test doubles (Stub, Fake, Mock) verified in tests/test_add_mark.py.",
        "Real file persistence verified via tmp_path in tests/test_storage.py.",
        "NOTES.md updated with fault injection evidence and double classifications.",
    ]
    for item in checklist_items:
        r_check = p_check.add_run(f"✓  {item}\n")
        r_check.font.name = "Calibri"
        r_check.font.size = Pt(10)
        r_check.font.color.rgb = RGBColor(16, 185, 129)

    doc.save(DOC_NAME)
    print(f"Report saved successfully as: {DOC_NAME}")

if __name__ == "__main__":
    build_report()
